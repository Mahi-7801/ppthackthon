import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = pptx.Presentation()
prs.slide_width = Inches(13.333)  # 16:9 Widescreen
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Ultra-Premium Color Palette (Modern Executive Light/Dark Theme)
DARK_BG = RGBColor(11, 19, 43)        # #0B132B Deep Navy
LIGHT_BG = RGBColor(248, 250, 252)    # #F8FAFC Slate 50
CARD_WHITE = RGBColor(255, 255, 255)
CARD_DARK = RGBColor(23, 37, 84)      # #172554
BORDER_GRAY = RGBColor(226, 232, 240) # #E2E8F0
PRIMARY_BLUE = RGBColor(37, 99, 235)  # #2563EB Royal Blue
CYAN_ACCENT = RGBColor(6, 182, 212)   # #06B6D4
SUCCESS_GREEN = RGBColor(16, 185, 129)# #10B981 Emerald
SUCCESS_DARK = RGBColor(5, 150, 105)  # #059669
DANGER_RED = RGBColor(239, 68, 68)    # #EF4444
DANGER_DARK = RGBColor(220, 38, 38)   # #DC2626
AMBER_GOLD = RGBColor(245, 158, 11)   # #F59E0B
PURPLE_ACC = RGBColor(124, 58, 237)   # #7C3AED
TEXT_MAIN = RGBColor(15, 23, 42)      # #0F172A Dark Slate
TEXT_MUTED = RGBColor(71, 85, 105)    # #475569 Slate
LIGHT_SLATE = RGBColor(148, 163, 184) # #94A3B8
WHITE = RGBColor(255, 255, 255)

def set_bg(slide, color=LIGHT_BG):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.color.rgb = color
    return bg

def add_header(slide, title, category="GOVERNMENT OF ANDHRA PRADESH • RTIH • APIS • NIC INNOVATION CHALLENGE 2026", is_dark=False):
    # Top Accent Bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY_BLUE
    bar.line.color.rgb = PRIMARY_BLUE
    
    # Category Tag
    tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.7), Inches(0.3))
    p_cat = tb_cat.text_frame.paragraphs[0]
    p_cat.text = category.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = PRIMARY_BLUE if not is_dark else CYAN_ACCENT
    
    # Title
    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(0.65))
    p_title = tb_title.text_frame.paragraphs[0]
    p_title.text = title
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_MAIN if not is_dark else WHITE

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 1: Title Slide (High-Impact Hero)
# ══════════════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(blank_layout)
set_bg(s1, DARK_BG)

# Glowing Background Shape
glow = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
glow.fill.solid()
glow.fill.fore_color.rgb = RGBColor(17, 24, 39)
glow.line.color.rgb = PRIMARY_BLUE
glow.line.width = Pt(2)

# Top Badge
pill = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.4), Inches(1.3), Inches(6.5), Inches(0.45))
pill.fill.solid()
pill.fill.fore_color.rgb = RGBColor(30, 58, 138)
pill.line.color.rgb = PRIMARY_BLUE
p_p = pill.text_frame.paragraphs[0]
p_p.alignment = PP_ALIGN.CENTER
p_p.text = "🏛️ AP GOVT • RTIH • APIS • NIC INNOVATION CHALLENGE 2026"
p_p.font.size = Pt(10.5)
p_p.font.bold = True
p_p.font.color.rgb = WHITE

# Main Title
tb1 = s1.shapes.add_textbox(Inches(1.4), Inches(1.9), Inches(10.5), Inches(2.2))
tf1 = tb1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "SecureSign: Type-C DSC Mobile Signing"
p.font.size = Pt(34)
p.font.bold = True
p.font.color.rgb = WHITE
p.space_after = Pt(8)

p2 = tf1.add_paragraph()
p2.text = "Direct, Hardware-Level Digital Signatures on Android Smartphones with Zero Key Leakage"
p2.font.size = Pt(16)
p2.font.italic = True
p2.font.color.rgb = CYAN_ACCENT
p2.space_after = Pt(12)

p3 = tf1.add_paragraph()
p3.text = "Lead Innovator: pmahi7801@gmail.com  •  Live API: https://hackthonapp-production.up.railway.app"
p3.font.size = Pt(12)
p3.font.color.rgb = LIGHT_SLATE

# 3 Bottom Stat Pills
stat_data = [
    ("⚡ < 800ms Latency", "Direct USB CCID Driver", PRIMARY_BLUE),
    ("🔒 Zero Key Extraction", "FIPS 140-2 Level 3 (CCA Rule 1)", SUCCESS_GREEN),
    ("⚖️ IT Act 2000 Sec 3A", "PAdES-LTV & RFC 3161 TSA", AMBER_GOLD)
]

for idx, (stitle, ssub, scolor) in enumerate(stat_data):
    sbox = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.4 + idx * 3.6), Inches(4.7), Inches(3.3), Inches(1.4))
    sbox.fill.solid()
    sbox.fill.fore_color.rgb = RGBColor(23, 37, 84)
    sbox.line.color.rgb = scolor
    sbox.line.width = Pt(1.5)
    
    stf = sbox.text_frame
    stf.word_wrap = True
    sp0 = stf.paragraphs[0]
    sp0.text = stitle
    sp0.font.size = Pt(13)
    sp0.font.bold = True
    sp0.font.color.rgb = scolor
    sp0.space_after = Pt(4)
    
    sp1 = stf.add_paragraph()
    sp1.text = ssub
    sp1.font.size = Pt(10)
    sp1.font.color.rgb = WHITE

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 2: Problem vs. Solution (Rich Side-by-Side Cards)
# ══════════════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(blank_layout)
set_bg(s2, LIGHT_BG)
add_header(s2, "The Desktop Bottleneck vs. The SecureSign Innovation", "CHALLENGE & VALUE PROPOSITION")

# Left: Problem Card
p_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.65), Inches(5.4))
p_card.fill.solid()
p_card.fill.fore_color.rgb = CARD_WHITE
p_card.line.color.rgb = DANGER_RED
p_card.line.width = Pt(1.5)

# Problem Header Banner
p_hdr = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.65), Inches(0.65))
p_hdr.fill.solid()
p_hdr.fill.fore_color.rgb = DANGER_RED
p_hdr.line.color.rgb = DANGER_RED
p_hp = p_hdr.text_frame.paragraphs[0]
p_hp.text = "❌ THE DESKTOP BOTTLENECK (CURRENT STATE)"
p_hp.font.size = Pt(12)
p_hp.font.bold = True
p_hp.font.color.rgb = WHITE
p_hp.alignment = PP_ALIGN.CENTER

# Problem Body
p_tb = s2.shapes.add_textbox(Inches(1.0), Inches(2.25), Inches(5.25), Inches(4.5))
p_tf = p_tb.text_frame
p_tf.word_wrap = True

prob_items = [
    ("Desktop-Only Confinement", "Officers are tied to desktop PCs with Windows to sign official files with USB DSC dongles."),
    ("Fragile Legacy Middleware", "Heavy reliance on unstable Java applets, browser plugins, and vendor desktop drivers (ePass/Watchdata tools)."),
    ("Zero Mobile Support", "Android lacks native smart card CCID drivers, preventing direct phone-based hardware signing."),
    ("Administrative Delays", "Cabinet files (e-Office), CFMS treasury bills, and citizen certificates (MeeSeva) stall when officers travel.")
]

for idx, (title, desc) in enumerate(prob_items):
    p0 = p_tf.add_paragraph() if idx > 0 else p_tf.paragraphs[0]
    p0.text = f"{idx+1}. {title}"
    p0.font.size = Pt(12.5)
    p0.font.bold = True
    p0.font.color.rgb = DANGER_DARK
    p0.space_after = Pt(2)
    
    p1 = p_tf.add_paragraph()
    p1.text = desc
    p1.font.size = Pt(10.5)
    p1.font.color.rgb = TEXT_MUTED
    p1.space_after = Pt(10)

# Right: Solution Card
s_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.88), Inches(1.5), Inches(5.65), Inches(5.4))
s_card.fill.solid()
s_card.fill.fore_color.rgb = CARD_WHITE
s_card.line.color.rgb = SUCCESS_GREEN
s_card.line.width = Pt(1.5)

# Solution Header Banner
s_hdr = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.88), Inches(1.5), Inches(5.65), Inches(0.65))
s_hdr.fill.solid()
s_hdr.fill.fore_color.rgb = SUCCESS_GREEN
s_hdr.line.color.rgb = SUCCESS_GREEN
s_hp = s_hdr.text_frame.paragraphs[0]
s_hp.text = "✅ THE SECURESIGN INNOVATION (OUR SOLUTION)"
s_hp.font.size = Pt(12)
s_hp.font.bold = True
s_hp.font.color.rgb = WHITE
s_hp.alignment = PP_ALIGN.CENTER

# Solution Body
s_tb = s2.shapes.add_textbox(Inches(7.08), Inches(2.25), Inches(5.25), Inches(4.5))
s_tf = s_tb.text_frame
s_tf.word_wrap = True

sol_items = [
    ("Direct USB Type-C CCID Driver", "Custom native Kotlin driver communicates directly with DSC dongles on Android without desktop middleware."),
    ("Zero Key Leakage (CCA Rule 1)", "Private key remains 100% secured inside the FIPS 140-2 Level 3 crypto chip. Key extraction is impossible."),
    ("On-Chip PIN Verification (Rule 2)", "Token PIN verified directly on-chip via VERIFY APDU; RAM is zeroized (0x00) immediately."),
    ("PAdES-LTV + RFC 3161 Standard", "Produces globally verifiable, tamper-evident signed PDFs with trusted timestamps in <3 seconds.")
]

for idx, (title, desc) in enumerate(sol_items):
    p0 = s_tf.add_paragraph() if idx > 0 else s_tf.paragraphs[0]
    p0.text = f"{idx+1}. {title}"
    p0.font.size = Pt(12.5)
    p0.font.bold = True
    p0.font.color.rgb = SUCCESS_DARK
    p0.space_after = Pt(2)
    
    p1 = s_tf.add_paragraph()
    p1.text = desc
    p1.font.size = Pt(10.5)
    p1.font.color.rgb = TEXT_MUTED
    p1.space_after = Pt(10)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 3: Visual Overview of the Diagram (Clean & Large)
# ══════════════════════════════════════════════════════════════════════════
s3 = prs.slides.add_slide(blank_layout)
set_bg(s3, LIGHT_BG)
add_header(s3, "System Architecture & Activity Diagram: Visual Overview", "4-PHASE CRYPTOGRAPHIC LIFECYCLE")

# Embed 3D Isometric Engineering Diagram Image
img_path = 'uploads/SecureSign_3D_Engine_Architecture_Blueprint.png'
if not os.path.exists(img_path):
    img_path = 'uploads/SecureSign_Executive_Workflow_Diagram.png'
if os.path.exists(img_path):
    s3.shapes.add_picture(img_path, Inches(0.8), Inches(1.45), Inches(11.733), Inches(5.6))

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 4: 4-Phase Step-by-Step Technical Flow (4 Beautiful Columns)
# ══════════════════════════════════════════════════════════════════════════
s4 = prs.slides.add_slide(blank_layout)
set_bg(s4, LIGHT_BG)
add_header(s4, "4-Phase Cryptographic Process Flow: Step-by-Step", "TECHNICAL ARCHITECTURE")

phases_spec = [
    ("PHASE 1", "01. INTAKE", "Mobile Device", PRIMARY_BLUE, [
        "User selects PDF in SecureSign Android app.",
        "On-device engine computes 32-byte SHA-256 hash.",
        "Original document never leaves unencrypted.",
        "Zero Document Exposure."
    ]),
    ("PHASE 2", "02. CCID", "USB Token", PURPLE_ACC, [
        "Android USB Host claims CCID Class 0x0B.",
        "Officer inputs 8-digit PIN on secure screen.",
        "Dispatches ISO 7816-4 VERIFY APDU to chip.",
        "RAM zeroized immediately (CCA Rule 2)."
    ]),
    ("PHASE 3", "03. RSA CHIP", "FIPS 140-2 L3", AMBER_GOLD, [
        "Token coprocessor executes RSA-2048 signing.",
        "Private Key NEVER leaves token (CCA Rule 1).",
        "Hardware locks after 3 failed PINs (Rule 4).",
        "Generates PKCS#1v1.5 signature."
    ]),
    ("PHASE 4", "04. LEGAL", "PAdES Delivery", SUCCESS_GREEN, [
        "Cloud backend injects RFC 3161 TSA timestamp.",
        "pdf-lib stamps AP Govt CCA Class-3 seal.",
        "1-Tap download to Adobe Acrobat (Zero OTP).",
        "100% Valid under IT Act Sec 3A."
    ])
]

for idx, (p_num, p_title, p_sub, p_color, p_bullets) in enumerate(phases_spec):
    left_pos = Inches(0.8 + idx * 2.98)
    
    # Base Card
    card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(1.5), Inches(2.78), Inches(5.4))
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_WHITE
    card.line.color.rgb = p_color
    card.line.width = Pt(1.5)
    
    # Header Strip
    hdr_s = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, Inches(1.5), Inches(2.78), Inches(0.75))
    hdr_s.fill.solid()
    hdr_s.fill.fore_color.rgb = p_color
    hdr_s.line.color.rgb = p_color
    
    htf = hdr_s.text_frame
    hp0 = htf.paragraphs[0]
    hp0.text = p_title
    hp0.font.size = Pt(12)
    hp0.font.bold = True
    hp0.font.color.rgb = WHITE
    hp0.alignment = PP_ALIGN.CENTER
    
    hp1 = htf.add_paragraph()
    hp1.text = p_sub
    hp1.font.size = Pt(9)
    hp1.font.color.rgb = RGBColor(240, 240, 240)
    hp1.alignment = PP_ALIGN.CENTER
    
    # Body
    tb_b = s4.shapes.add_textbox(left_pos + Inches(0.15), Inches(2.35), Inches(2.48), Inches(4.4))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    
    for b_idx, bullet in enumerate(p_bullets):
        bp = tf_b.add_paragraph() if b_idx > 0 else tf_b.paragraphs[0]
        bp.text = "• " + bullet
        bp.font.size = Pt(10.5)
        bp.font.color.rgb = TEXT_MAIN
        bp.space_after = Pt(10)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 5: CCA India Regulatory Compliance Matrix
# ══════════════════════════════════════════════════════════════════════════
s5 = prs.slides.add_slide(blank_layout)
set_bg(s5, LIGHT_BG)
add_header(s5, "100% CCA India Regulatory Compliance Matrix", "LEGAL & STATUTORY VERIFICATION")

cca_rules = [
    ("Rule 1: Private Key Isolation", "Private key must NEVER leave hardware token under any circumstance.", "On-chip RSA-2048 signing inside FIPS 140-2 Level 3 Secure Element.", "100% PASS"),
    ("Rule 2: Token PIN Verification", "PIN must be verified directly on-chip without memory caching.", "ISO 7816-4 VERIFY APDU; PIN memory zeroized (0x00) immediately.", "100% PASS"),
    ("Rule 3: Signature Standard", "PAdES / CAdES standard with trusted Time Stamping Authority (TSA).", "ETSI EN 319 142-1 (PAdES-LTV) with embedded RFC 3161 TSA token.", "100% PASS"),
    ("Rule 4: Hardware Locking", "Enforce hardware-level brute-force retry limits.", "Hardware token locks automatically after 3 consecutive wrong PINs (SW1=0x63).", "100% PASS"),
    ("Rule 5: Audit Trail", "Maintain immutable, tamper-evident audit logs.", "Cryptographic audit trail with signer ID, SHA-256 hash, IP & timestamp.", "100% PASS")
]

for idx, (r_name, r_mandate, r_impl, r_status) in enumerate(cca_rules):
    top_pos = Inches(1.5 + idx * 1.08)
    
    # Row Card
    rcard = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.733), Inches(0.95))
    rcard.fill.solid()
    rcard.fill.fore_color.rgb = CARD_WHITE
    rcard.line.color.rgb = BORDER_GRAY
    rcard.line.width = Pt(1)
    
    # Status Pill
    spill = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.0), top_pos + Inches(0.25), Inches(1.3), Inches(0.45))
    spill.fill.solid()
    spill.fill.fore_color.rgb = RGBColor(220, 252, 231)
    spill.line.color.rgb = SUCCESS_GREEN
    spp = spill.text_frame.paragraphs[0]
    spp.alignment = PP_ALIGN.CENTER
    spp.text = "✅ " + r_status
    spp.font.size = Pt(9.5)
    spp.font.bold = True
    spp.font.color.rgb = SUCCESS_DARK
    
    # Text
    rtb = s5.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.08), Inches(9.8), Inches(0.8))
    rtf = rtb.text_frame
    rtf.word_wrap = True
    
    rp0 = rtf.paragraphs[0]
    rp0.text = r_name
    rp0.font.size = Pt(12)
    rp0.font.bold = True
    rp0.font.color.rgb = PRIMARY_BLUE
    rp0.space_after = Pt(2)
    
    rp1 = rtf.add_paragraph()
    rp1.text = f"Mandate: {r_mandate}  ➔  Implementation: {r_impl}"
    rp1.font.size = Pt(10)
    rp1.font.color.rgb = TEXT_MUTED

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 6: Real-World AP Governance Impact (4 High-Impact Grid Cards)
# ══════════════════════════════════════════════════════════════════════════
s6 = prs.slides.add_slide(blank_layout)
set_bg(s6, LIGHT_BG)
add_header(s6, "Transforming Andhra Pradesh Digital Governance", "REAL-WORLD DEPLOYMENT")

impact_grid = [
    ("📁 e-Office & Secretariat Approvals", "Ministers and IAS officers can sign and approve urgent cabinet notes, policy memos, and Government Orders (G.O.s) directly on mobile phones while traveling.", PRIMARY_BLUE),
    ("💰 CFMS Treasury Bill Clearances", "Drawing and Disbursing Officers (DDOs) across 26 districts can clear treasury bills, salary vouchers, and vendor payments on smartphones in seconds.", SUCCESS_GREEN),
    ("📜 MeeSeva Citizen Services", "Revenue Officers (Tahsildars/MROs) can issue digitally signed caste, income, and Land Title (ROR-1B) certificates directly in the field.", PURPLE_ACC),
    ("🏢 AP e-Procurement & Tenders", "Instant digital contractor signature verification for state infrastructure tenders without requiring desktop workstations.", AMBER_GOLD)
]

for idx, (i_title, i_desc, i_color) in enumerate(impact_grid):
    r_idx = idx // 2
    c_idx = idx % 2
    
    l_pos = Inches(0.8 + c_idx * 5.95)
    t_pos = Inches(1.5 + r_idx * 2.7)
    
    icard = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l_pos, t_pos, Inches(5.78), Inches(2.45))
    icard.fill.solid()
    icard.fill.fore_color.rgb = CARD_WHITE
    icard.line.color.rgb = i_color
    icard.line.width = Pt(1.5)
    
    itb = s6.shapes.add_textbox(l_pos + Inches(0.2), t_pos + Inches(0.2), Inches(5.38), Inches(2.05))
    itf = itb.text_frame
    itf.word_wrap = True
    
    ip0 = itf.paragraphs[0]
    ip0.text = i_title
    ip0.font.size = Pt(13)
    ip0.font.bold = True
    ip0.font.color.rgb = i_color
    ip0.space_after = Pt(6)
    
    ip1 = itf.add_paragraph()
    ip1.text = i_desc
    ip1.font.size = Pt(11)
    ip1.font.color.rgb = TEXT_MUTED
    ip1.space_after = Pt(4)

# ══════════════════════════════════════════════════════════════════════════
# SLIDE 7: Live Verification Benchmarks & Submission Links
# ══════════════════════════════════════════════════════════════════════════
s7 = prs.slides.add_slide(blank_layout)
set_bg(s7, DARK_BG)
add_header(s7, "Live Performance Benchmarks & Submission Links", "VERIFICATION SUMMARY", is_dark=True)

# 4 Performance Metric Boxes
metrics = [
    ("0.42s", "Signing Latency", "Sub-second processing across REST APIs", PRIMARY_BLUE),
    ("13 / 13", "Endpoints 200 OK", "All master verification tests passing", SUCCESS_GREEN),
    ("100%", "CCA Valid", "Adobe Acrobat green verification ribbon", CYAN_ACCENT),
    ("0 Keys", "Zero Key Leakage", "FIPS 140-2 Level 3 Secure Element", AMBER_GOLD)
]

for idx, (m_val, m_lbl, m_sub, m_col) in enumerate(metrics):
    m_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + idx * 2.98), Inches(1.5), Inches(2.78), Inches(2.2))
    m_box.fill.solid()
    m_box.fill.fore_color.rgb = RGBColor(17, 24, 39)
    m_box.line.color.rgb = m_col
    m_box.line.width = Pt(1.5)
    
    mtf = m_box.text_frame
    mtf.word_wrap = True
    
    mp0 = mtf.paragraphs[0]
    mp0.text = m_val
    mp0.font.size = Pt(28)
    mp0.font.bold = True
    mp0.font.color.rgb = m_col
    mp0.alignment = PP_ALIGN.CENTER
    mp0.space_after = Pt(2)
    
    mp1 = mtf.add_paragraph()
    mp1.text = m_lbl
    mp1.font.size = Pt(12)
    mp1.font.bold = True
    mp1.font.color.rgb = WHITE
    mp1.alignment = PP_ALIGN.CENTER
    mp1.space_after = Pt(2)
    
    mp2 = mtf.add_paragraph()
    mp2.text = m_sub
    mp2.font.size = Pt(9)
    mp2.font.color.rgb = LIGHT_SLATE
    mp2.alignment = PP_ALIGN.CENTER

# Bottom Submission Links Card
link_card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.0), Inches(11.733), Inches(2.9))
link_card.fill.solid()
link_card.fill.fore_color.rgb = RGBColor(17, 24, 39)
link_card.line.color.rgb = PRIMARY_BLUE
link_card.line.width = Pt(1.5)

ltb = s7.shapes.add_textbox(Inches(1.1), Inches(4.2), Inches(11.1), Inches(2.5))
ltf = ltb.text_frame
ltf.word_wrap = True

lp0 = ltf.paragraphs[0]
lp0.text = "🚀 OFFICIAL SUBMISSION RESOURCES & CONTACT"
lp0.font.size = Pt(14)
lp0.font.bold = True
lp0.font.color.rgb = PRIMARY_BLUE
lp0.space_after = Pt(8)

links_data = [
    ("Live Production API:", "https://hackthonapp-production.up.railway.app"),
    ("GitHub Source Code:", "https://github.com/mahankalikornepati2-netizen/hackthonapp"),
    ("Android APK (Build #4384cd86):", "https://expo.dev/accounts/mahibujjipapas-team/projects/dsc-mobile-signing/builds/4384cd86-8e73-42f3-ad81-033d7bbb9d2c"),
    ("Lead Innovator Contact:", "pmahi7801@gmail.com | AP Government Innovation Challenge 2026")
]

for l_label, l_url in links_data:
    lp = ltf.add_paragraph()
    r1 = lp.add_run()
    r1.text = "• " + l_label + " "
    r1.font.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = WHITE
    
    r2 = lp.add_run()
    r2.text = l_url
    r2.font.size = Pt(11)
    r2.font.color.rgb = CYAN_ACCENT
    lp.space_after = Pt(4)

os.makedirs('uploads', exist_ok=True)
out_pptx1 = 'uploads/SecureSign_Executive_Pitch_Deck.pptx'
out_pptx2 = 'SecureSign_Executive_Pitch_Deck.pptx'

prs.save(out_pptx1)
prs.save(out_pptx2)

print(f"Successfully generated World-Class PowerPoint Presentation at:")
print(f"1. {out_pptx1}")
print(f"2. {out_pptx2}")
