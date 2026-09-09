import os
from pptx import Presentation

pptx_files = [f for f in os.listdir('.') if f.endswith('.pptx')]
for f in pptx_files:
    try:
        prs = Presentation(f)
        print(f"\n==================== {f} (slides: {len(prs.slides)}) ====================")
        for i, s in enumerate(prs.slides):
            texts = []
            for shape in s.shapes:
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t:
                        texts.append(t.replace('\n', ' '))
            print(f"Slide {i+1}: {' // '.join(texts[:2])[:120]}")
    except Exception as e:
        print(f"{f}: {e}")
