import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

os.makedirs('uploads', exist_ok=True)

# ══════════════════════════════════════════════════════════════════════
# 1. BUILD OFFICIAL UNSIGNED TEST PDF (REPORTLAB)
# ══════════════════════════════════════════════════════════════════════
pdf_path = 'uploads/SecureSign_Evaluation_Test_Document_Unsigned.pdf'
pdf_root = 'SecureSign_Evaluation_Test_Document_Unsigned.pdf'

doc_pdf = SimpleDocTemplate(
    pdf_path,
    pagesize=letter,
    rightMargin=45, leftMargin=45, topMargin=40, bottomMargin=40
)

styles = getSampleStyleSheet()

header_style = ParagraphStyle(
    'GovHeader',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=13,
    leading=17,
    textColor=colors.HexColor('#0B132B'),
    alignment=1,
    spaceAfter=2
)

sub_header = ParagraphStyle(
    'GovSub',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=10,
    leading=14,
    textColor=colors.HexColor('#1E293B'),
    alignment=1,
    spaceAfter=6
)

order_style = ParagraphStyle(
    'OrderStyle',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=11,
    leading=15,
    textColor=colors.HexColor('#2563EB'),
    alignment=1,
    spaceAfter=12
)

body_style = ParagraphStyle(
    'GovBody',
    parent=styles['Normal'],
    fontName='Helvetica',
    fontSize=9.5,
    leading=14,
    textColor=colors.HexColor('#1E293B'),
    spaceAfter=8
)

body_bold = ParagraphStyle(
    'GovBodyBold',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=9.5,
    leading=14,
    textColor=colors.HexColor('#0B132B'),
    spaceAfter=8
)

story = []

story.append(Paragraph("GOVERNMENT OF ANDHRA PRADESH", header_style))
story.append(Paragraph("INFORMATION TECHNOLOGY, ELECTRONICS & COMMUNICATIONS (IT&C) DEPARTMENT<br/>ANDHRA PRADESH SECRETARIAT, VELAGAPUDI, AMARAVATI", sub_header))
story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#0B132B'), spaceAfter=8))

story.append(Paragraph("<u>G.O. Ms. No. 104 &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; Dated: 28-08-2026</u>", order_style))

abstract_text = (
    "<b><u>ABSTRACT</u></b><br/>"
    "<b>E-Governance — Secure Mobile Digital Signing Framework (SecureSign) — Implementation across all State Secretariats, "
    "Directorates, and District Collectorates — Administrative Sanction Accorded — Orders — Issued.</b>"
)
story.append(Paragraph(abstract_text, body_style))
story.append(Spacer(1, 4))

ref_text = (
    "<b><u>Read the following:</u></b><br/>"
    "1. G.O. Rt. No. 42, ITE&C Dept., dated 14-01-2026.<br/>"
    "2. Recommendations of the High-Level Committee on Digital Governance & PKI Security, dated 10-06-2026.<br/>"
    "3. Guidelines of the Controller of Certifying Authorities (CCA), Ministry of Electronics & IT, Government of India."
)
story.append(Paragraph(ref_text, body_style))
story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#94A3B8'), spaceAfter=8))

story.append(Paragraph("<b><u>ORDER:</u></b>", body_bold))

order_p1 = (
    "1. The Government of Andhra Pradesh is committed to delivering responsive, agile, and paperless e-governance services. "
    "Under current workflows, administrative approvals and financial sanctions requiring Class-3 Digital Signature Certificates (DSC) "
    "have been tied to desktop personal computers, leading to procedural delays whenever competent authorities are on tour or in field inspections."
)
story.append(Paragraph(order_p1, body_style))

order_p2 = (
    "2. Following rigorous evaluation under the SecureSign Innovation Challenge 2026, administrative sanction is hereby accorded "
    "for state-wide deployment of the <b>SecureSign Mobile Digital Signing Solution</b>. The solution enables direct connection of standard "
    "USB Type-C and OTG hardware DSC dongles to mobile smartphones and tablets, facilitating instantaneous, CCA-compliant cryptographic signing "
    "under Sections 3 & 3A of the Indian Information Technology Act, 2000."
)
story.append(Paragraph(order_p2, body_style))

order_p3 = (
    "3. All Special Chief Secretaries, Principal Secretaries, District Collectors, and Heads of Departments are instructed to adopt "
    "the mobile digital signing framework for e-Office file approvals, MeeSeva service issuances, and CFMS bill clearances."
)
story.append(Paragraph(order_p3, body_style))
story.append(Spacer(1, 10))

# Signoff and Signature Box Placeholder
sign_box = [
    [
        Paragraph("<b>(BY ORDER AND IN THE NAME OF THE GOVERNOR OF ANDHRA PRADESH)</b><br/><br/><b>DR. K. VIJAYANAND, IAS</b><br/>Special Chief Secretary to Government", body_style),
        Paragraph(
            "<font color='#94A3B8'><b>[ OFFICIAL SIGNATURE TARGET BOX ]</b><br/>"
            "Status: <i>Pending Digital Signature</i><br/>"
            "Use SecureSign App to Sign This Document<br/>"
            "Hardware: USB Type-C DSC Token / Sandbox</font>",
            ParagraphStyle('SigBox', parent=body_style, fontName='Helvetica', fontSize=8, leading=11, alignment=1)
        )
    ]
]
t_sign = Table(sign_box, colWidths=[3.8*inch, 3.2*inch])
t_sign.setStyle(TableStyle([
    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ('BOX', (1, 0), (1, 0), 1.5, colors.HexColor('#2563EB')),
    ('BACKGROUND', (1, 0), (1, 0), colors.HexColor('#F8FAFC')),
    ('TOPPADDING', (1, 0), (1, 0), 12),
    ('BOTTOMPADDING', (1, 0), (1, 0), 12),
    ('LEFTPADDING', (1, 0), (1, 0), 10),
    ('RIGHTPADDING', (1, 0), (1, 0), 10),
]))
story.append(t_sign)

story.append(Spacer(1, 14))
story.append(Paragraph(
    "<b>To:</b><br/>"
    "The Commissioner of e-Governance & IT, Andhra Pradesh.<br/>"
    "All District Collectors and District Magistrates, Government of Andhra Pradesh.<br/>"
    "Copy to: The Director General, National Informatics Centre (NIC), New Delhi.<br/>"
    "Copy to: PS to Hon'ble Chief Minister, Government of Andhra Pradesh.",
    ParagraphStyle('ToStyle', parent=body_style, fontSize=8.5, leading=12, textColor=colors.HexColor('#475569'))
))

doc_pdf.build(story)

# Copy to root
import shutil
shutil.copyfile(pdf_path, pdf_root)

# ══════════════════════════════════════════════════════════════════════
# 2. BUILD MATCHING OFFICIAL UNSIGNED TEST WORD DOC (DOCX)
# ══════════════════════════════════════════════════════════════════════
docx_path = 'uploads/SecureSign_Evaluation_Test_Document_Unsigned.docx'
docx_root = 'SecureSign_Evaluation_Test_Document_Unsigned.docx'

doc_word = docx.Document()
for section in doc_word.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

p_w1 = doc_word.add_paragraph()
p_w1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_w1 = p_w1.add_run("GOVERNMENT OF ANDHRA PRADESH\n")
r_w1.bold = True
r_w1.font.size = Pt(13)
r_w1.font.color.rgb = RGBColor(11, 19, 43)

r_w2 = p_w1.add_run("INFORMATION TECHNOLOGY, ELECTRONICS & COMMUNICATIONS DEPARTMENT\nANDHRA PRADESH SECRETARIAT, VELAGAPUDI, AMARAVATI")
r_w2.bold = True
r_w2.font.size = Pt(10)
r_w2.font.color.rgb = RGBColor(30, 41, 59)

p_w2 = doc_word.add_paragraph()
p_w2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_w3 = p_w2.add_run("G.O. Ms. No. 104                                                          Dated: 28-08-2026")
r_w3.bold = True
r_w3.underline = True
r_w3.font.size = Pt(11)
r_w3.font.color.rgb = RGBColor(37, 99, 235)

p_w3 = doc_word.add_paragraph()
r_w4 = p_w3.add_run("ABSTRACT\n")
r_w4.bold = True
r_w4.underline = True
r_w5 = p_w3.add_run("E-Governance — Secure Mobile Digital Signing Framework (SecureSign) — Implementation across all State Secretariats, Directorates, and District Collectorates — Administrative Sanction Accorded — Orders — Issued.")
r_w5.bold = True

p_w4 = doc_word.add_paragraph()
r_w6 = p_w4.add_run("ORDER:\n")
r_w6.bold = True
r_w6.underline = True

doc_word.add_paragraph(
    "1. The Government of Andhra Pradesh is committed to delivering responsive, agile, and paperless e-governance services. "
    "Under current workflows, administrative approvals and financial sanctions requiring Class-3 Digital Signature Certificates (DSC) "
    "have been tied to desktop personal computers, leading to procedural delays whenever competent authorities are on tour or in field inspections."
)

doc_word.add_paragraph(
    "2. Following rigorous evaluation under the SecureSign Innovation Challenge 2026, administrative sanction is hereby accorded "
    "for state-wide deployment of the SecureSign Mobile Digital Signing Solution. The solution enables direct connection of standard "
    "USB Type-C and OTG hardware DSC dongles to mobile smartphones and tablets, facilitating instantaneous, CCA-compliant cryptographic signing "
    "under Sections 3 & 3A of the Indian Information Technology Act, 2000."
)

doc_word.add_paragraph(
    "3. All Special Chief Secretaries, Principal Secretaries, District Collectors, and Heads of Departments are instructed to adopt "
    "the mobile digital signing framework for e-Office file approvals, MeeSeva service issuances, and CFMS bill clearances."
)

# Signature table in Word
sig_t = doc_word.add_table(rows=1, cols=2)
sig_t.alignment = WD_TABLE_ALIGNMENT.CENTER
c1 = sig_t.rows[0].cells[0]
c2 = sig_t.rows[0].cells[1]
c1.width = Inches(4.0)
c2.width = Inches(2.8)

p_c1 = c1.paragraphs[0]
r_c1 = p_c1.add_run("(BY ORDER AND IN THE NAME OF THE GOVERNOR OF ANDHRA PRADESH)\n\nDR. K. VIJAYANAND, IAS\nSpecial Chief Secretary to Government")
r_c1.bold = True
r_c1.font.size = Pt(9.5)

p_c2 = c2.paragraphs[0]
p_c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_c2 = p_c2.add_run("[ OFFICIAL SIGNATURE TARGET BOX ]\nStatus: Pending Digital Signature\nUse SecureSign App to Sign This Document\nHardware: USB Type-C DSC Token / Sandbox")
r_c2.font.size = Pt(8.5)
r_c2.font.italic = True
r_c2.font.color.rgb = RGBColor(100, 116, 139)

doc_word.save(docx_path)
doc_word.save(docx_root)

print(f"Generated Unsigned Test Documents successfully:\n- {pdf_path}\n- {docx_path}")
