import os
import matplotlib.pyplot as plt
import matplotlib.patches as patches
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

os.makedirs('uploads', exist_ok=True)

# ── 1. Create High-Resolution UML Activity Diagram Graphic ──
fig, ax = plt.subplots(figsize=(16, 11), dpi=300)
ax.set_xlim(0, 16)
ax.set_ylim(0, 11)
ax.axis('off')

# Style Constants
BOX_COLOR = '#FFFFFF'
BOX_EDGE = '#0F172A'
TEXT_COLOR = '#0F172A'
DECISION_COLOR = '#FFFFFF'
DECISION_EDGE = '#007AFF'
ARROW_COLOR = '#0F172A'
FAIL_BOX_COLOR = '#FEF2F2'
FAIL_EDGE = '#EF4444'
SUCCESS_BOX_COLOR = '#F0FDF4'
SUCCESS_EDGE = '#10B981'

def draw_activity_box(ax, x, y, width, height, text, fill=BOX_COLOR, edge=BOX_EDGE, bold=False):
    rect = patches.FancyBboxPatch(
        (x, y), width, height,
        boxstyle="round,pad=0.08,rounding_size=0.15",
        facecolor=fill, edgecolor=edge, linewidth=1.8, zorder=3
    )
    ax.add_patch(rect)
    ax.text(
        x + width / 2, y + height / 2, text,
        ha='center', va='center', fontsize=9.5, fontweight='bold' if bold else 'normal',
        color=TEXT_COLOR, zorder=4, linespacing=1.2
    )

def draw_diamond(ax, cx, cy, size, text):
    pts = [
        [cx, cy + size],
        [cx + size * 1.3, cy],
        [cx, cy - size],
        [cx - size * 1.3, cy]
    ]
    diamond = patches.Polygon(pts, facecolor=DECISION_COLOR, edgecolor=DECISION_EDGE, linewidth=1.8, zorder=3)
    ax.add_patch(diamond)
    ax.text(
        cx, cy, text,
        ha='center', va='center', fontsize=8.5, fontweight='bold',
        color='#007AFF', zorder=4
    )

def draw_start_node(ax, cx, cy):
    circle = patches.Circle((cx, cy), radius=0.25, facecolor='#0F172A', edgecolor='#0F172A', zorder=4)
    ax.add_patch(circle)

def draw_end_node(ax, cx, cy):
    c_out = patches.Circle((cx, cy), radius=0.3, facecolor='none', edgecolor='#0F172A', linewidth=2, zorder=4)
    c_in = patches.Circle((cx, cy), radius=0.18, facecolor='#0F172A', edgecolor='#0F172A', zorder=4)
    ax.add_patch(c_out)
    ax.add_patch(c_in)

def draw_arrow(ax, start, end, label=None, label_offset=(0, 0)):
    ax.annotate(
        '', xy=end, xytext=start,
        arrowprops=dict(facecolor=ARROW_COLOR, edgecolor=ARROW_COLOR, arrowstyle='->,head_width=0.25,head_length=0.35', lw=1.5),
        zorder=2
    )
    if label:
        lx = (start[0] + end[0]) / 2 + label_offset[0]
        ly = (start[1] + end[1]) / 2 + label_offset[1]
        ax.text(lx, ly, label, fontsize=8.5, fontweight='bold', color='#1E293B', bbox=dict(facecolor='white', edgecolor='none', pad=1))

# Title Header
ax.text(8.0, 10.4, "SECURESIGN — UML ACTIVITY DIAGRAM", ha='center', va='center', fontsize=18, fontweight='bold', color='#0B132B')
ax.text(8.0, 9.95, "End-to-End Type-C DSC Hardware Signing & Verification Process Flow (CCA India Compliant)", ha='center', va='center', fontsize=10.5, fontstyle='italic', color='#64748B')

# ── Draw Nodes ──

# Start Node
draw_start_node(ax, 0.8, 8.8)

# Row 1: Document Selection & Hash
draw_activity_box(ax, 1.6, 8.35, 2.4, 0.9, "Pick PDF Document\n(Android Storage)")
draw_arrow(ax, (1.05, 8.8), (1.6, 8.8))

draw_activity_box(ax, 4.6, 8.35, 2.5, 0.9, "Compute SHA-256 Hash\n(On-Device Crypto)")
draw_arrow(ax, (4.0, 8.8), (4.6, 8.8))

# Decision 1: Token Connected?
draw_diamond(ax, 8.2, 8.8, 0.45, "Token\nConnected?")
draw_arrow(ax, (7.1, 8.8), (7.6, 8.8))

# Branch 1 No: Show Token Error
draw_activity_box(ax, 7.0, 6.7, 2.4, 0.8, "Display 'Plug Type-C\nDSC Token' Alert", fill=FAIL_BOX_COLOR, edge=FAIL_EDGE)
draw_arrow(ax, (8.2, 8.35), (8.2, 7.5), label="No", label_offset=(0.25, 0))

draw_end_node(ax, 8.2, 5.7)
draw_arrow(ax, (8.2, 6.7), (8.2, 6.0))

# Branch 1 Yes: Prompt PIN Entry
draw_activity_box(ax, 9.9, 8.35, 2.2, 0.9, "Prompt Token PIN\n(Secure Screen)")
draw_arrow(ax, (8.8, 8.8), (9.9, 8.8), label="Yes", label_offset=(0, 0.2))

# Row 2: PIN Verification & Hardware Signing
draw_activity_box(ax, 12.8, 8.35, 2.5, 0.9, "Send VERIFY APDU\n[00 20 00 81 PIN]")
draw_arrow(ax, (12.1, 8.8), (12.8, 8.8))

# Move Down from Row 1 to Row 2
draw_arrow(ax, (14.05, 8.35), (14.05, 5.4))

# Decision 2: PIN Valid?
draw_diamond(ax, 14.05, 4.8, 0.48, "PIN\nValid?")

# Branch 2 No: Brute Force Guard
draw_activity_box(ax, 10.2, 4.35, 2.5, 0.9, "Retries >= 3 ?\nLock Hardware Token", fill=FAIL_BOX_COLOR, edge=FAIL_EDGE)
draw_arrow(ax, (13.4, 4.8), (12.7, 4.8), label="No", label_offset=(0, 0.2))

draw_end_node(ax, 9.2, 4.8)
draw_arrow(ax, (10.2, 4.8), (9.5, 4.8))

# Loop back for retries < 3
ax.annotate(
    '', xy=(11.0, 8.35), xytext=(11.45, 5.25),
    arrowprops=dict(facecolor='#64748B', edgecolor='#64748B', arrowstyle='->', lw=1.2, linestyle='dashed')
)
ax.text(11.6, 6.8, "Retry (<3)", fontsize=7.5, color='#64748B', fontweight='bold')

# Branch 2 Yes: Move Down to Row 3 (Signing Flow)
draw_arrow(ax, (14.05, 4.3), (14.05, 2.9), label="Yes", label_offset=(0.25, 0))

# Row 3: Signing, Timestamp, PAdES, Stamping, Download
draw_activity_box(ax, 12.6, 2.0, 2.9, 0.9, "On-Chip RSA Signing\n(PSO:SIGN APDU 00 2A)")

draw_activity_box(ax, 9.3, 2.0, 2.8, 0.9, "Request RFC 3161 TSA\n(Trusted Timestamp)")
draw_arrow(ax, (12.6, 2.45), (12.1, 2.45))

draw_activity_box(ax, 6.0, 2.0, 2.8, 0.9, "Assemble PAdES-LTV\n& Stamp Visible Seal", fill=SUCCESS_BOX_COLOR, edge=SUCCESS_EDGE)
draw_arrow(ax, (9.3, 2.45), (8.8, 2.45))

draw_activity_box(ax, 2.7, 2.0, 2.8, 0.9, "1-Tap PDF Download\n& Open in Acrobat", fill=SUCCESS_BOX_COLOR, edge=SUCCESS_EDGE, bold=True)
draw_arrow(ax, (6.0, 2.45), (5.5, 2.45))

# Final Success Node
draw_end_node(ax, 1.2, 2.45)
draw_arrow(ax, (2.7, 2.45), (1.5, 2.45))
ax.text(1.2, 1.8, "Signature Valid\n(IT Act Sec 3A)", ha='center', fontsize=8, fontweight='bold', color='#10B981')

# Bottom Banner Box
legend_rect = patches.FancyBboxPatch(
    (0.8, 0.3), 14.4, 0.8,
    boxstyle="round,pad=0.05,rounding_size=0.1",
    facecolor='#F8FAFC', edgecolor='#CBD5E1', linewidth=1.2
)
ax.add_patch(legend_rect)
ax.text(8.0, 0.7, "LEGEND:  ⬤ Initial Start  |  [ Rounded Box ] Action State  |  ◇ Decision Node  |  ◉ Final End State  |  🔴 Brute-Force Guard", ha='center', va='center', fontsize=8.5, fontweight='bold', color='#334155')

plt.tight_layout()
png_path1 = 'uploads/SecureSign_Activity_Diagram.png'
png_path2 = 'SecureSign_Activity_Diagram.png'
plt.savefig(png_path1, dpi=300, bbox_inches='tight')
plt.savefig(png_path2, dpi=300, bbox_inches='tight')
plt.close()

print(f"Generated PNG Activity Diagram: {png_path1}")

# ── 2. Create Microsoft Word Document (.docx) Embedding this Diagram ──
doc = docx.Document()

# Set standard page margins
for section in doc.sections:
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

# Title Page / Header
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title_p.add_run('SECURESIGN: UML ACTIVITY DIAGRAM SPECIFICATION\n')
title_run.bold = True
title_run.font.size = Pt(20)
title_run.font.color.rgb = RGBColor(11, 19, 43)

sub_run = title_p.add_run('Complete Process Flowchart & Decision Branching for Type-C DSC Hardware Signing\n')
sub_run.font.size = Pt(12)
sub_run.font.italic = True
sub_run.font.color.rgb = RGBColor(71, 85, 105)

meta_run = title_p.add_run('Government of Andhra Pradesh • RTIH • APIS • NIC Innovation Challenge 2026\nLead Contact: pmahi7801@gmail.com | Live API: https://hackthonapp-production.up.railway.app\n')
meta_run.font.size = Pt(9.5)
meta_run.font.color.rgb = RGBColor(0, 122, 255)

doc.add_paragraph('─' * 75)

# Insert High-Resolution Activity Diagram Image into Word
p_img = doc.add_paragraph()
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img.paragraph_format.space_before = Pt(8)
p_img.paragraph_format.space_after = Pt(12)
doc.add_picture(png_path1, width=Inches(6.8))

doc.add_paragraph('')

# Activity Description Table
p_tbl_title = doc.add_paragraph()
r_tt = p_tbl_title.add_run('ACTIVITY DIAGRAM NODES & DECISION SPECIFICATION')
r_tt.bold = True
r_tt.font.size = Pt(13)
r_tt.font.color.rgb = RGBColor(11, 19, 43)

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

chdrs = table.rows[0].cells
chdrs[0].width = Inches(1.8)
chdrs[1].width = Inches(2.2)
chdrs[2].width = Inches(2.8)
chdrs[0].text = 'Activity Node / Decision'
chdrs[1].text = 'Technical Action / APDU'
chdrs[2].text = 'Security & Compliance Guarantee'

for c in chdrs:
    set_cell_background(c, '0B132B')
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)

nodes_data = [
    ('1. Pick PDF & SHA-256', 'Android DocumentPicker & Crypto.digestStringAsync()', 'Zero Document Exposure: Only mathematical 32-byte hash is generated on-device.'),
    ('2. Decision: Token Connected?', 'android.hardware.usb.UsbManager detection of CCID Class 0x0B', 'Direct USB Host Interface: Claims USB endpoint 0x82/0x02 without desktop drivers.'),
    ('3. Enter PIN & VERIFY APDU', 'ISO 7816-4: [00 20 00 81 <PIN_BYTES>]', 'CCA Rule 2: PIN verified on hardware chip; RAM is zeroized (0x00) immediately.'),
    ('4. Decision: PIN Valid?', 'Token checks PIN against security chip retry counter', 'CCA Rule 4: Hardware locks token after 3 failed attempts (Brute-Force Guard).'),
    ('5. On-Chip RSA Signing', 'ISO 7816-4: [00 2A 9E 9A <DigestInfo + Hash>]', 'CCA Rule 1: Private key NEVER leaves FIPS 140-2 Level 3 crypto chip.'),
    ('6. RFC 3161 TSA Timestamp', 'HTTP POST timestamp query to trusted TSA', 'CCA Rule 3: Cryptographic proof of signing time embedded in signature dictionary.'),
    ('7. Assemble PAdES & Seal', 'pdf-lib embeds PAdES-LTV container & visible seal', 'Visual + Cryptographic Integrity: Document locked under IT Act 2000 Section 3A.'),
    ('8. Download & Acrobat View', '1-Tap open in Adobe Acrobat (Zero OTP)', 'Global Admissibility: Displays Green Checkmark "Signed & all signatures valid".')
]

for node, action, sec in nodes_data:
    row = table.add_row().cells
    row[0].width = Inches(1.8)
    row[1].width = Inches(2.2)
    row[2].width = Inches(2.8)
    
    row[0].text = node
    row[1].text = action
    row[2].text = sec
    
    set_cell_background(row[0], 'F0F4F8')
    set_cell_background(row[1], 'FFFFFF')
    set_cell_background(row[2], 'FFFFFF')
    
    for c in row:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.size = Pt(8.5)

doc.add_paragraph('')

# Verification Links
p_links = doc.add_paragraph()
p_links.add_run('Official Submission & Verification Links:\n').bold = True
p_links.add_run('• Production API (Railway): https://hackthonapp-production.up.railway.app\n')
p_links.add_run('• GitHub Source Code: https://github.com/mahankalikornepati2-netizen/hackthonapp\n')
p_links.add_run('• Android APK Build #4384cd86: https://expo.dev/accounts/mahibujjipapas-team/projects/dsc-mobile-signing/builds/4384cd86-8e73-42f3-ad81-033d7bbb9d2c\n')
p_links.add_run('• Official Google Form: https://forms.gle/TJDYkF6feKFrywsd7\n')

docx_path1 = 'uploads/SecureSign_Activity_Diagram.docx'
docx_path2 = 'SecureSign_Activity_Diagram.docx'

doc.save(docx_path1)
doc.save(docx_path2)

print('Successfully generated Word Document with Activity Diagram at:')
print('1.', docx_path1)
print('2.', docx_path2)
