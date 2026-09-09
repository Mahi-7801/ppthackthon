from pptx import Presentation

prs = Presentation('DSC_Mobile_Signing_VoneDigital_Deck.pptx')
for i, slide in enumerate(prs.slides):
    print(f"\n==================== Slide {i+1} ====================")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text.strip():
                    print("  -", p.text.strip())
