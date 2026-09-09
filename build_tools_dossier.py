import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

os.makedirs('uploads', exist_ok=True)
doc = docx.Document()

# Page Margins
for section in doc.sections:
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(11, 19, 43) # Deep Navy
    return p

# ── Title & Executive Header ──
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title_p.add_run('SECURESIGN: TOOLS, FRAMEWORKS & COMPONENTS SPECIFICATION\n')
title_run.bold = True
title_run.font.size = Pt(18)
title_run.font.color.rgb = RGBColor(11, 19, 43)

sub_run = title_p.add_run('Comprehensive Technology Stack, Libraries, SDKs, Protocols & Architectural Dependencies\n')
sub_run.font.size = Pt(11)
sub_run.font.italic = True
sub_run.font.color.rgb = RGBColor(71, 85, 105)

meta_run = title_p.add_run('Government of Andhra Pradesh • RTIH • APIS • NIC Innovation Challenge 2026\nLead Contact: pmahi7801@gmail.com | Production API: https://hackthonapp-production.up.railway.app\n')
meta_run.font.size = Pt(9.5)
meta_run.font.color.rgb = RGBColor(37, 99, 235)

doc.add_paragraph('─' * 75)

# ── Section 1: Technology Stack Summary Table ──
add_heading_1(doc, '1. Comprehensive Technology Stack & Component Inventory')

stack_table = doc.add_table(rows=1, cols=3)
stack_table.alignment = WD_TABLE_ALIGNMENT.CENTER
stack_table.autofit = False

shdr = stack_table.rows[0].cells
shdr[0].width = Inches(1.8)
shdr[1].width = Inches(2.4)
shdr[2].width = Inches(2.7)
shdr[0].text = 'Layer / Tier'
shdr[1].text = 'Frameworks, SDKs & Libraries'
shdr[2].text = 'Purpose & Implementation Details'

for c in shdr:
    set_cell_background(c, '0B132B')
    set_cell_margins(c, top=80, bottom=80, left=120, right=120)
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(255, 255, 255)

stack_rows = [
    ('1. Mobile Frontend', '• React Native 0.74+\n• TypeScript 5.x\n• Expo SDK 51 & EAS Build\n• React Navigation v6\n• expo-document-picker\n• expo-crypto (SHA-256)', 'Cross-platform mobile UI/UX, native file picking, on-device SHA-256 hashing, PIN entry screen, 1-tap download & sharing with zero OTP delays.'),
    ('2. Native Android Driver', '• Kotlin 1.9+\n• android.hardware.usb\n• UsbManager & UsbDeviceConnection\n• CcidTransport.kt\n• P11Wrapper.kt\n• DSCSigningModule.kt', 'Direct USB Type-C CCID host driver claiming USB Class 0x0B Smart Card endpoints (0x82/0x02), managing APDUs without desktop middleware.'),
    ('3. Hardware Crypto Chip', '• ISO/IEC 7816-4 APDU\n• PKCS#11 v2.40 Token API\n• PKCS#15 Structure\n• FIPS 140-2 Level 3 Coprocessor\n• CC EAL 5+ Secure Element', 'On-chip RSA-2048 signing, on-chip PIN verification (VERIFY APDU), hardware brute-force defense (3 failed attempts lock token). Zero key leakage.'),
    ('4. Cloud PAdES Backend', '• Node.js v20 LTS\n• Express.js REST API\n• pdf-lib Engine\n• Native Node.js Crypto Engine\n• RFC 3161 TSA Client', 'PAdES-LTV (ETSI TS 102 778) container generation, trusted RFC 3161 timestamping, stamping official AP Govt visible seal on Page 2 of user PDF.'),
    ('5. Cloud DB & Hosting', '• Railway Cloud Platform (HTTPS)\n• Supabase PostgreSQL 15\n• Row-Level Security (RLS)\n• JWT Authentication (HS256)', 'High-availability REST API (<0.5s response time), user authentication, tamper-evident cryptographic audit logs, document session mapping.')
]

for layer, libs, purpose in stack_rows:
    row = stack_table.add_row().cells
    row[0].width = Inches(1.8)
    row[1].width = Inches(2.4)
    row[2].width = Inches(2.7)
    row[0].text = layer
    row[1].text = libs
    row[2].text = purpose
    
    set_cell_background(row[0], 'F8FAFC')
    set_cell_background(row[1], 'FFFFFF')
    set_cell_background(row[2], 'FFFFFF')
    set_cell_margins(row[0], top=60, bottom=60, left=100, right=100)
    set_cell_margins(row[1], top=60, bottom=60, left=100, right=100)
    set_cell_margins(row[2], top=60, bottom=60, left=100, right=100)
    
    for c in row:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.size = Pt(8.5)

doc.add_paragraph('')

# ── Section 2: Hardware & Dongle Compatibility Matrix ──
add_heading_1(doc, '2. Hardware Tokens, Standards & Certifying Authorities')

compat_table = doc.add_table(rows=1, cols=3)
compat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
compat_table.autofit = False

chdr = compat_table.rows[0].cells
chdr[0].width = Inches(2.0)
chdr[1].width = Inches(2.4)
chdr[2].width = Inches(2.5)
chdr[0].text = 'Category'
chdr[1].text = 'Supported Standards & Vendors'
chdr[2].text = 'Compliance & Verification'

for c in chdr:
    set_cell_background(c, '0B132B')
    set_cell_margins(c, top=80, bottom=80, left=120, right=120)
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(255, 255, 255)

compat_rows = [
    ('Hardware Dongles', '• Feitian ePass2003 / Auto\n• Watchdata PROXKey / TrustKey\n• Gemalto SafeNet IDPrime\n• mToken CryptoID / HyperSecu', 'FIPS 140-2 Level 3 / Common Criteria EAL 5+ validated.'),
    ('Certifying Authorities (CAs)', '• e-Mudhra • Capricorn\n• VSign • Sify\n• (n)Code Solutions • Pantasign', 'CCA Class-3 Signing Certificates verified against Indian Root CA.'),
    ('Cryptographic Algorithms', '• SHA-256, SHA-384, SHA-512\n• RSA 2048-bit / 4096-bit\n• ECDSA (NIST P-256)', 'FIPS 180-4 / FIPS 186-4 cryptographic standard compliant.'),
    ('Digital Signature Standard', '• PAdES-LTV (ETSI EN 319 142-1)\n• RFC 3161 / RFC 5816 TSA\n• ISO 32000-1 PDF Standard', 'Section 3 & 3A of the Indian Information Technology Act 2000.')
]

for cat, ven, comp in compat_rows:
    row = compat_table.add_row().cells
    row[0].width = Inches(2.0)
    row[1].width = Inches(2.4)
    row[2].width = Inches(2.5)
    row[0].text = cat
    row[1].text = ven
    row[2].text = comp
    
    set_cell_background(row[0], 'F8FAFC')
    set_cell_background(row[1], 'FFFFFF')
    set_cell_background(row[2], 'FFFFFF')
    set_cell_margins(row[0], top=60, bottom=60, left=100, right=100)
    set_cell_margins(row[1], top=60, bottom=60, left=100, right=100)
    set_cell_margins(row[2], top=60, bottom=60, left=100, right=100)
    
    for c in row:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.size = Pt(8.5)

doc.add_paragraph('')

# ── Section 3: Official Verification & Repository Links ──
add_heading_1(doc, '3. Production Verification & Repository References')

p_links = doc.add_paragraph()
p_links.add_run('• Production REST API: ').bold = True
p_links.add_run('https://hackthonapp-production.up.railway.app\n')
p_links.add_run('• GitHub Source Code: ').bold = True
p_links.add_run('https://github.com/mahankalikornepati2-netizen/hackthonapp\n')
p_links.add_run('• Android Standalone APK (Build #4384cd86): ').bold = True
p_links.add_run('https://expo.dev/accounts/mahibujjipapas-team/projects/dsc-mobile-signing/builds/4384cd86-8e73-42f3-ad81-033d7bbb9d2c\n')
p_links.add_run('• Lead Innovator Email: ').bold = True
p_links.add_run('pmahi7801@gmail.com\n')

out_docx1 = 'uploads/SecureSign_Tools_Frameworks_and_Components_Specification.docx'
out_docx2 = 'SecureSign_Tools_Frameworks_and_Components_Specification.docx'

doc.save(out_docx1)
doc.save(out_docx2)

print(f"Generated Tools Specification Dossier at:")
print(f"1. {out_docx1}")
print(f"2. {out_docx2}")
